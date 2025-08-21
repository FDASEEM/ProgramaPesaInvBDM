import tkinter as tk
from tkinter import ttk, messagebox, filedialog, simpledialog
import serial
import serial.tools.list_ports
import threading
import time
import json
from datetime import datetime
import queue
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

class ModernPesaUSBMonitor:
    def __init__(self, root):
        self.root = root
        self.root.title("Pesa Inv Bosques - v2.0")
        
        # tamaño fijo y centrado en pantalla
        width, height = 1300, 900
        self.root.update_idletasks()
        screen_w = self.root.winfo_screenwidth()
        screen_h = self.root.winfo_screenheight()
        x = (screen_w - width) // 2
        y = (screen_h - height) // 2
        self.root.geometry(f"{width}x{height}+{x}+{y}")
        self.root.configure(bg='#f0f0f0')
        
        # Configurar tema moderno
        self.setup_modern_theme()
        
        # Variables
        self.serial_connection = None
        self.is_connected = False
        self.is_monitoring = False
        self.data_queue = queue.Queue()
        self.peso_actual = 0.0
        self.historial_pesos = []
        self.historial_tiempos = []
        
        # Barcode / lector de pistola
        self.barcode_buffer = ""
        self.last_key_time = 0.0
        self.scanner_inter_char_delay = 0.1  # segundos, para resetear buffer si demora mucho
        self.barcode_map = {}  # barcode -> {"rut":..., "name":...}
        self.mapping_filename = os.path.join(os.path.expanduser("~"), "barcode_map.json")
        self.current_worker = None  # dict {"rut":..., "name":...}
        
        # Variables para el proceso de carga
        self.peso_inicial = None
        self.peso_carga_actual = None
        self.proceso_activo = False
        self.bandejas_retiradas = []
        
        # Configurar interfaz moderna
        self.setup_modern_ui()
        
        # Hilo para procesar datos
        self.monitor_thread = None
        self.root.after(100, self.process_queue)
        
        # Bind teclas para lector de pistola
        self.root.bind("<Key>", self.on_key_press)
        self.root.focus_set()
        
        # Animaciones
        self.peso_color_animation = 0
    
    def setup_modern_theme(self):
        """Configura un tema moderno para la aplicación"""
        style = ttk.Style()
        
        # Tema base
        style.theme_use('clam')
        
        # Colores modernos
        colors = {
            'bg': '#2c3e50',
            'fg': '#ecf0f1',
            'select_bg': '#3498db',
            'select_fg': '#ffffff',
            'button_bg': '#3498db',
            'button_hover': '#2980b9',
            'success': '#27ae60',
            'danger': '#e74c3c',
            'warning': '#f39c12'
        }
        
        # Configurar estilos
        style.configure('Modern.TFrame', background='#34495e', relief='flat')
        style.configure('Card.TFrame', background='#ffffff', relief='solid', borderwidth=1)
        style.configure('Modern.TLabel', background='#34495e', foreground='#ecf0f1', font=('Segoe UI', 10))
        style.configure('Card.TLabel', background='#ffffff', foreground='#2c3e50', font=('Segoe UI', 10))
        style.configure('Title.TLabel', background='#34495e', foreground='#ecf0f1', font=('Segoe UI', 14, 'bold'))
        
        # Botones modernos
        style.configure('Modern.TButton', 
                       font=('Segoe UI', 10),
                       borderwidth=0,
                       focuscolor='none',
                       background='#3498db',
                       foreground='white')
        style.map('Modern.TButton',
                 background=[('active', '#2980b9'), ('pressed', '#21618c')])
        
        style.configure('Success.TButton', 
                       background='#27ae60',
                       foreground='white')
        style.map('Success.TButton',
                 background=[('active', '#229954'), ('pressed', '#1e8449')])
        
        style.configure('Danger.TButton', 
                       background='#e74c3c',
                       foreground='white')
        style.map('Danger.TButton',
                 background=[('active', '#c0392b'), ('pressed', '#a93226')])
    
    def setup_modern_ui(self):
        # Frame principal con fondo moderno
        main_frame = tk.Frame(self.root, bg='#34495e', padx=20, pady=20)
        main_frame.pack(fill='both', expand=True)
        
        # Header con título
        header_frame = tk.Frame(main_frame, bg='#34495e', height=80)
        header_frame.pack(fill='x', pady=(0, 20))
        header_frame.pack_propagate(False)
        
        title_label = tk.Label(header_frame, 
                              text="🔗 Sistema de Pesaje de Bandejas", 
                              bg='#34495e', 
                              fg='#ecf0f1',
                              font=('Segoe UI', 18, 'bold'))
        title_label.pack(pady=20)
        
        # Container principal con grid
        content_frame = tk.Frame(main_frame, bg='#34495e')
        content_frame.pack(fill='both', expand=True)
        
        # Tarjeta de conexión
        self.create_connection_card(content_frame)
        
        # Tarjeta de peso principal
        self.create_weight_display_card(content_frame)
        
        # Tarjeta de controles
        self.create_controls_card(content_frame)
        
        # Tarjeta de proceso de carga
        self.create_carga_card(content_frame)
        
        # Tarjeta de historial
        self.create_history_card(content_frame)
        
        # Tarjeta de mediciones guardadas
        self.create_saved_measurements_card(content_frame)
        
        # Configurar grid
        content_frame.grid_columnconfigure(0, weight=1)
        content_frame.grid_columnconfigure(1, weight=1)
        content_frame.grid_rowconfigure(2, weight=1)
        content_frame.grid_rowconfigure(3, weight=1)
        
        # Inicializar
        self.actualizar_puertos()
    
    def create_connection_card(self, parent):
        """Crea la tarjeta de conexión moderna"""
        card = tk.Frame(parent, bg='#ffffff', relief='solid', bd=1)
        card.grid(row=0, column=0, columnspan=2, sticky='ew', pady=(0, 15), padx=5)
        
        # Header de la tarjeta
        header = tk.Frame(card, bg='#3498db', height=40)
        header.pack(fill='x')
        header.pack_propagate(False)
        
        tk.Label(header, text="🔌 Conexión", bg='#3498db', fg='white',
                font=('Segoe UI', 12, 'bold')).pack(side='left', padx=15, pady=10)
        
        # Contenido
        content = tk.Frame(card, bg='#ffffff', padx=20, pady=15)
        content.pack(fill='x')
        
        # Fila de controles
        controls_row = tk.Frame(content, bg='#ffffff')
        controls_row.pack(fill='x')
        
        tk.Label(controls_row, text="Puerto:", bg='#ffffff', fg='#2c3e50',
                font=('Segoe UI', 10)).pack(side='left', padx=(0, 10))
        
        self.puerto_combo = ttk.Combobox(controls_row, state="readonly", width=15,
                                        font=('Segoe UI', 10))
        self.puerto_combo.pack(side='left', padx=(0, 10))
        
        refresh_btn = tk.Button(controls_row, text="🔄 Actualizar", 
                               command=self.actualizar_puertos,
                               bg='#95a5a6', fg='white', relief='flat',
                               font=('Segoe UI', 9), cursor='hand2')
        refresh_btn.pack(side='left', padx=(0, 15))
        
        # Botón para cargar mapeos de barcode -> RUT
        load_map_btn = tk.Button(controls_row, text="📥 Cargar RUTs", 
                                 command=self.load_barcode_mapping,
                                 bg='#8e44ad', fg='white', relief='flat',
                                 font=('Segoe UI', 9), cursor='hand2')
        load_map_btn.pack(side='left', padx=(0, 10))
        
        self.btn_conectar = tk.Button(controls_row, text="🔗 Conectar",
                                     command=self.toggle_connection,
                                     bg='#27ae60', fg='white', relief='flat',
                                     font=('Segoe UI', 10, 'bold'), 
                                     cursor='hand2', width=12)
        self.btn_conectar.pack(side='left')
        
        # Status
        status_frame = tk.Frame(content, bg='#ffffff')
        status_frame.pack(fill='x', pady=(10, 0))
        
        tk.Label(status_frame, text="Estado:", bg='#ffffff', fg='#2c3e50',
                font=('Segoe UI', 10)).pack(side='left')
        
        self.status_label = tk.Label(status_frame, text="⭕ Desconectado", 
                                    bg='#ffffff', fg='#e74c3c',
                                    font=('Segoe UI', 10, 'bold'))
        self.status_label.pack(side='left', padx=(5, 0))
        
        # Label para el trabajador actual detectado por barcode
        worker_frame = tk.Frame(content, bg='#ffffff')
        worker_frame.pack(fill='x', pady=(8, 0))
        tk.Label(worker_frame, text="Trabajador:", bg='#ffffff', fg='#2c3e50',
                font=('Segoe UI', 10)).pack(side='left')
        self.worker_label = tk.Label(worker_frame, text="— Ninguno —", bg='#ffffff', fg='#2c3e50',
                                     font=('Segoe UI', 10, 'bold'))
        self.worker_label.pack(side='left', padx=(8, 0))
    
    def create_weight_display_card(self, parent):
        """Crea la tarjeta del display de peso principal"""
        card = tk.Frame(parent, bg='#ffffff', relief='solid', bd=1)
        card.grid(row=1, column=0, sticky='ew', pady=(0, 15), padx=(5, 10))
        
        # Header
        header = tk.Frame(card, bg='#2c3e50', height=40)
        header.pack(fill='x')
        header.pack_propagate(False)
        
        tk.Label(header, text="⚖️ Peso Actual", bg='#2c3e50', fg='white',
                font=('Segoe UI', 12, 'bold')).pack(side='left', padx=15, pady=10)
        
        # Display grande del peso
        display_frame = tk.Frame(card, bg='#ffffff', height=120)
        display_frame.pack(fill='both', expand=True)
        display_frame.pack_propagate(False)
        
        self.peso_var = tk.StringVar(value="0.00")
        self.peso_display = tk.Label(display_frame, textvariable=self.peso_var,
                                    bg='#ffffff', fg='#2c3e50',
                                    font=('Segoe UI', 32, 'bold'))
        self.peso_display.pack(expand=True)
        
        # Unidad
        tk.Label(display_frame, text="kilogramos", bg='#ffffff', fg='#7f8c8d',
                font=('Segoe UI', 12)).pack()
    
    def create_controls_card(self, parent):
        """Crea la tarjeta de controles"""
        card = tk.Frame(parent, bg='#ffffff', relief='solid', bd=1)
        card.grid(row=1, column=1, sticky='new', pady=(0, 15), padx=(10, 5))
        
        # Header
        header = tk.Frame(card, bg='#9b59b6', height=40)
        header.pack(fill='x')
        header.pack_propagate(False)
        
        tk.Label(header, text="🎛️ Controles", bg='#9b59b6', fg='white',
                font=('Segoe UI', 12, 'bold')).pack(side='left', padx=15, pady=10)
        
        # Contenido
        content = tk.Frame(card, bg='#ffffff', padx=20, pady=20)
        content.pack(fill='both', expand=True)
        
        # Botones con iconos y estilo moderno
        self.btn_tara = tk.Button(content, text="🎯 Hacer Tara",
                                 command=self.hacer_tara,
                                 bg='#f39c12', fg='white', relief='flat',
                                 font=('Segoe UI', 11, 'bold'),
                                 cursor='hand2', state='disabled',
                                 width=15, pady=8)
        self.btn_tara.pack(pady=(0, 10), fill='x')
        
        self.btn_guardar = tk.Button(content, text="💾 Guardar Peso",
                                   command=self.guardar_peso,
                                   bg='#27ae60', fg='white', relief='flat',
                                   font=('Segoe UI', 11, 'bold'),
                                   cursor='hand2', state='disabled',
                                   width=15, pady=8)
        self.btn_guardar.pack(pady=(0, 10), fill='x')
        
        self.btn_exportar = tk.Button(content, text="📤 Exportar Datos",
                                     command=self.exportar_datos,
                                     bg='#3498db', fg='white', relief='flat',
                                     font=('Segoe UI', 11, 'bold'),
                                     cursor='hand2', width=15, pady=8)
        self.btn_exportar.pack(fill='x')
    
    def create_carga_card(self, parent):
        """Crea la tarjeta para el proceso de carga"""
        card = tk.Frame(parent, bg='#ffffff', relief='solid', bd=1)
        card.grid(row=2, column=0, sticky='ew', pady=(0, 15), padx=(5, 10))
        
        # Header
        header = tk.Frame(card, bg='#e67e22', height=40)
        header.pack(fill='x')
        header.pack_propagate(False)
        
        tk.Label(header, text="📦 Proceso de Carga", bg='#e67e22', fg='white',
                font=('Segoe UI', 12, 'bold')).pack(side='left', padx=15, pady=10)
        
        # Contenido
        content = tk.Frame(card, bg='#ffffff', padx=20, pady=15)
        content.pack(fill='x')
        
        # Peso inicial
        peso_frame = tk.Frame(content, bg='#ffffff')
        peso_frame.pack(fill='x', pady=(0, 10))
        
        tk.Label(peso_frame, text="Peso Inicial:", bg='#ffffff', fg='#2c3e50',
                font=('Segoe UI', 10)).pack(side='left')
        
        self.peso_inicial_var = tk.StringVar(value="0.00 kg")
        tk.Label(peso_frame, textvariable=self.peso_inicial_var, bg='#ffffff', fg='#2c3e50',
                font=('Segoe UI', 10, 'bold')).pack(side='left', padx=(5, 0))
        
        # Peso actual de carga
        peso_actual_frame = tk.Frame(content, bg='#ffffff')
        peso_actual_frame.pack(fill='x', pady=(0, 10))
        
        tk.Label(peso_actual_frame, text="Peso Carga Actual:", bg='#ffffff', fg='#2c3e50',
                font=('Segoe UI', 10)).pack(side='left')
        
        self.peso_carga_actual_var = tk.StringVar(value="0.00 kg")
        tk.Label(peso_actual_frame, textvariable=self.peso_carga_actual_var, bg='#ffffff', fg='#2c3e50',
                font=('Segoe UI', 10, 'bold')).pack(side='left', padx=(5, 0))
        
        # Botones de control de proceso
        btn_frame = tk.Frame(content, bg='#ffffff')
        btn_frame.pack(fill='x', pady=(10, 0))
        
        self.btn_iniciar_proceso = tk.Button(btn_frame, text="▶️ Iniciar Proceso",
                                           command=self.iniciar_proceso,
                                           bg='#27ae60', fg='white', relief='flat',
                                           font=('Segoe UI', 10, 'bold'),
                                           cursor='hand2', state='disabled')
        self.btn_iniciar_proceso.pack(side='left', padx=(0, 10))
        
        self.btn_finalizar_proceso = tk.Button(btn_frame, text="⏹️ Finalizar Proceso",
                                             command=self.finalizar_proceso,
                                             bg='#e74c3c', fg='white', relief='flat',
                                             font=('Segoe UI', 10, 'bold'),
                                             cursor='hand2', state='disabled')
        self.btn_finalizar_proceso.pack(side='left')
        
        # Estado del proceso
        estado_frame = tk.Frame(content, bg='#ffffff')
        estado_frame.pack(fill='x', pady=(10, 0))
        
        tk.Label(estado_frame, text="Estado Proceso:", bg='#ffffff', fg='#2c3e50',
                font=('Segoe UI', 10)).pack(side='left')
        
        self.estado_proceso_var = tk.StringVar(value="Inactivo")
        tk.Label(estado_frame, textvariable=self.estado_proceso_var, bg='#ffffff', fg='#e74c3c',
                font=('Segoe UI', 10, 'bold')).pack(side='left', padx=(5, 0))
    
    def create_history_card(self, parent):
        """Crea la tarjeta de historial moderno"""
        card = tk.Frame(parent, bg='#ffffff', relief='solid', bd=1)
        card.grid(row=3, column=0, sticky='nsew', pady=(0, 15), padx=(5, 10))
        
        # Header
        header = tk.Frame(card, bg='#1abc9c', height=40)
        header.pack(fill='x')
        header.pack_propagate(False)
        
        tk.Label(header, text="📊 Últimas Lecturas", bg='#1abc9c', fg='white',
                font=('Segoe UI', 12, 'bold')).pack(side='left', padx=15, pady=10)
        
        # Contenido
        content_frame = tk.Frame(card, bg='#ffffff', padx=15, pady=15)
        content_frame.pack(fill='both', expand=True)
        
        # Listbox con estilo moderno
        self.historial_listbox = tk.Listbox(content_frame, 
                                           font=('Consolas', 10),
                                           bg='#f8f9fa',
                                           fg='#2c3e50',
                                           selectbackground='#3498db',
                                           selectforeground='white',
                                           relief='flat',
                                           borderwidth=1)
        
        scrollbar_hist = ttk.Scrollbar(content_frame, orient="vertical", 
                                      command=self.historial_listbox.yview)
        self.historial_listbox.configure(yscrollcommand=scrollbar_hist.set)
        
        self.historial_listbox.pack(side='left', fill='both', expand=True)
        scrollbar_hist.pack(side='right', fill='y')
    
    def create_saved_measurements_card(self, parent):
        """Crea la tarjeta de mediciones guardadas"""
        card = tk.Frame(parent, bg='#ffffff', relief='solid', bd=1)
        card.grid(row=3, column=1, sticky='nsew', pady=(0, 15), padx=(10, 5))
        
        # Header
        header = tk.Frame(card, bg='#e67e22', height=40)
        header.pack(fill='x')
        header.pack_propagate(False)
        
        tk.Label(header, text="💾 Bandejas Retiradas", bg='#e67e22', fg='white',
                font=('Segoe UI', 12, 'bold')).pack(side='left', padx=15, pady=10)
        
        # Contenido
        content_frame = tk.Frame(card, bg='#ffffff', padx=15, pady=15)
        content_frame.pack(fill='both', expand=True)
        
        # Treeview con estilo moderno
        columns = ("Timestamp", "RUT", "Nombre", "Peso Bandeja", "Peso Carga Antes", "Peso Carga Después")
        self.tree = ttk.Treeview(content_frame, columns=columns, show="headings", height=10)
        self.tree.heading("Timestamp", text="📅 Fecha/Hora")
        self.tree.heading("RUT", text="🆔 RUT")
        self.tree.heading("Nombre", text="👤 Nombre")
        self.tree.heading("Peso Bandeja", text="⚖️ Peso Bandeja (kg)")
        self.tree.heading("Peso Carga Antes", text="📦 Peso Antes (kg)")
        self.tree.heading("Peso Carga Después", text="📦 Peso Después (kg)")
        
        self.tree.column("Timestamp", width=140, anchor='center')
        self.tree.column("RUT", width=100, anchor='center')
        self.tree.column("Nombre", width=120, anchor='w')
        self.tree.column("Peso Bandeja", width=100, anchor='center')
        self.tree.column("Peso Carga Antes", width=100, anchor='center')
        self.tree.column("Peso Carga Después", width=100, anchor='center')
        
        # Configurar colores alternos en las filas
        self.tree.tag_configure('oddrow', background='#f8f9fa')
        self.tree.tag_configure('evenrow', background='#ffffff')
        
        scrollbar_tree = ttk.Scrollbar(content_frame, orient="vertical", 
                                      command=self.tree.yview)
        self.tree.configure(yscrollcommand=scrollbar_tree.set)
        
        self.tree.pack(side='left', fill='both', expand=True)
        scrollbar_tree.pack(side='right', fill='y')
    
    def add_hover_effects(self):
        """Agrega efectos hover a los botones"""
        def on_enter(event, original_color, hover_color):
            event.widget.config(bg=hover_color)
        
        def on_leave(event, original_color):
            event.widget.config(bg=original_color)
        
        # Aplicar efectos a botones principales
        buttons = [
            (self.btn_conectar, '#27ae60', '#229954'),
            (self.btn_tara, '#f39c12', '#e67e22'),
            (self.btn_guardar, '#27ae60', '#229954'),
            (self.btn_exportar, '#3498db', '#2980b9'),
            (self.btn_iniciar_proceso, '#27ae60', '#229954'),
            (self.btn_finalizar_proceso, '#e74c3c', '#c0392b')
        ]
        
        for btn, original, hover in buttons:
            btn.bind("<Enter>", lambda e, o=original, h=hover: on_enter(e, o, h))
            btn.bind("<Leave>", lambda e, o=original: on_leave(e, o))
    
    def actualizar_puertos(self):
        """Actualiza la lista de puertos COM disponibles"""
        puertos = [port.device for port in serial.tools.list_ports.comports()]
        self.puerto_combo['values'] = puertos
        if puertos:
            self.puerto_combo.set(puertos[0])
    
    def toggle_connection(self):
        """Conectar o desconectar de la pesa"""
        if not self.is_connected:
            self.conectar()
        else:
            self.desconectar()
    
    def conectar(self):
        """Conecta con la pesa USB"""
        puerto = self.puerto_combo.get()
        if not puerto:
            messagebox.showerror("Error", "Selecciona un puerto COM")
            return
        
        try:
            self.serial_connection = serial.Serial(
                port=puerto,
                baudrate=9600,
                bytesize=serial.EIGHTBITS,
                parity=serial.PARITY_NONE,
                stopbits=serial.STOPBITS_ONE,
                timeout=1
            )
            
            self.is_connected = True
            self.is_monitoring = True
            
            # Actualizar UI con estilo
            self.btn_conectar.config(text="🔌 Desconectar", bg='#e74c3c')
            self.status_label.config(text="✅ Conectado", fg='#27ae60')
            self.btn_tara.config(state="normal")
            self.btn_guardar.config(state="normal")
            self.btn_iniciar_proceso.config(state="normal")
            self.puerto_combo.config(state="disabled")
            
            # Iniciar monitoreo
            self.monitor_thread = threading.Thread(target=self.monitor_peso, daemon=True)
            self.monitor_thread.start()
            
            # Mostrar notificación de éxito
            self.show_success_notification(f"Conectado al puerto {puerto}")
            
        except serial.SerialException as e:
            messagebox.showerror("Error de Conexión", f"No se pudo conectar: {e}")
    
    def desconectar(self):
        """Desconecta de la pesa"""
        self.is_monitoring = False
        self.is_connected = False
        
        if self.serial_connection:
            self.serial_connection.close()
            self.serial_connection = None
        
        # Actualizar UI
        self.btn_conectar.config(text="🔗 Conectar", bg='#27ae60')
        self.status_label.config(text="⭕ Desconectado", fg='#e74c3c')
        self.btn_tara.config(state="disabled")
        self.btn_guardar.config(state="disabled")
        self.btn_iniciar_proceso.config(state="disabled")
        self.btn_finalizar_proceso.config(state="disabled")
        self.puerto_combo.config(state="readonly")
        
        # Si hay un proceso activo, finalizarlo
        if self.proceso_activo:
            self.finalizar_proceso()
    
    def show_success_notification(self, message):
        """Muestra una notificación de éxito temporal centrada sobre la ventana principal"""
        notification = tk.Toplevel(self.root)
        notification.title("✅ Éxito")
        notification.configure(bg='#27ae60')
        notification.overrideredirect(True)
        notification.attributes("-topmost", True)
        notification.transient(self.root)

        label = tk.Label(notification, text="✅ " + message,
                         bg='#27ae60', fg='white',
                         font=('Segoe UI', 11, 'bold'), padx=20, pady=10)
        label.pack()

        # Forzar cálculo de tamaño y centrar sobre la ventana principal
        notification.update_idletasks()
        notif_w = notification.winfo_width()
        notif_h = notification.winfo_height()

        root_x = self.root.winfo_rootx()
        root_y = self.root.winfo_rooty()
        root_w = self.root.winfo_width()
        root_h = self.root.winfo_height()

        x = root_x + (root_w - notif_w) // 2
        y = root_y + (root_h - notif_h) // 2
        notification.geometry(f"{notif_w}x{notif_h}+{x}+{y}")

        # Auto-cerrar después de 2 segundos
        self.root.after(2000, notification.destroy)
    
    def monitor_peso(self):
        """Hilo que lee continuamente el peso de la pesa"""
        while self.is_monitoring:
            try:
                if self.serial_connection and self.serial_connection.in_waiting > 0:
                    data = self.serial_connection.readline().decode('utf-8').strip()
                    peso = self.parsear_peso(data)
                    if peso is not None:
                        self.data_queue.put(("peso", peso))
                
                time.sleep(0.1)
                
            except Exception as e:
                self.data_queue.put(("error", str(e)))
                break
    
    def parsear_peso(self, data):
        """Parsea los datos recibidos de la pesa"""
        try:
            if "WT:" in data:
                peso_str = data.split("WT:")[1].split("kg")[0].strip()
                return float(peso_str)
            elif data.replace(".", "").replace("-", "").isdigit():
                return float(data)
            else:
                print(f"Formato no reconocido: {data}")
                return None
        except (ValueError, IndexError):
            return None
    
    def process_queue(self):
        """Procesa los datos del hilo de monitoreo"""
        try:
            while True:
                item = self.data_queue.get_nowait()
                event_type, data = item
                
                if event_type == "peso":
                    self.actualizar_peso(data)
                elif event_type == "error":
                    messagebox.showerror("Error de Comunicación", data)
                    self.desconectar()
                
        except queue.Empty:
            pass
        
        self.root.after(100, self.process_queue)
    
    def actualizar_peso(self, peso):
        """Actualiza el display de peso con animación de color"""
        self.peso_actual = peso
        
        # Si hay un proceso activo, actualizar el peso de carga actual
        if self.proceso_activo:
            self.peso_carga_actual = peso
            self.peso_carga_actual_var.set(f"{peso:.2f} kg")
        
        self.peso_var.set(f"{peso:.2f}")
        
        # Animación de color para el peso
        self.animate_weight_display()
        
        # Agregar al historial
        timestamp = datetime.now().strftime("%H:%M:%S")
        entrada = f"{timestamp}  →  {peso:8.2f} kg"
        
        self.historial_listbox.insert(0, entrada)
        
        if self.historial_listbox.size() > 50:
            self.historial_listbox.delete(49)
        
        self.historial_listbox.see(0)
    
    def animate_weight_display(self):
        """Anima el color del display de peso"""
        colors = ['#2c3e50', '#3498db', '#2c3e50']
        current_color = colors[self.peso_color_animation % len(colors)]
        self.peso_display.config(fg=current_color)
        self.peso_color_animation += 1
        
        # Resetear después de un ciclo completo
        if self.peso_color_animation >= len(colors):
            self.peso_color_animation = 0
    
    def hacer_tara(self):
        """Envía comando de tara a la pesa"""
        if self.serial_connection:
            try:
                self.serial_connection.write(b'T\r\n')
                self.show_success_notification("Comando de tara enviado")
            except Exception as e:
                messagebox.showerror("Error", f"Error al hacer tara: {e}")
    
    def guardar_peso(self):
        """Guarda el peso actual en la lista con estilo alternado e información del trabajador"""
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        item_count = len(self.tree.get_children())
        tag = 'evenrow' if item_count % 2 == 0 else 'oddrow'
        
        rut = self.current_worker.get("rut") if self.current_worker else ""
        name = self.current_worker.get("name") if self.current_worker else ""
        
        self.tree.insert("", "end", values=(timestamp, rut, name, f"{self.peso_actual:.2f}", "", ""), tags=(tag,))
        
        # Auto-scroll al final
        children = self.tree.get_children()
        if children:
            self.tree.see(children[-1])
    
    def exportar_datos(self):
        """Exporta las mediciones guardadas a un archivo Excel"""
        if len(self.tree.get_children()) == 0:
            messagebox.showwarning("Sin Datos", "No hay mediciones para exportar")
            return
        
        filename = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
            title="Guardar datos como..."
        )
        
        if filename:
            try:
                # Crear un libro de trabajo de Excel
                wb = Workbook()
                ws = wb.active
                ws.title = "Registros de Pesaje"
                
                # Encabezados
                headers = ["Fecha/Hora", "RUT", "Nombre", "Peso Bandeja (kg)", "Peso Carga Antes (kg)", "Peso Carga Después (kg)"]
                for col, header in enumerate(headers, 1):
                    cell = ws.cell(row=1, column=col, value=header)
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
                    cell.alignment = Alignment(horizontal="center")
                
                # Datos
                for row, item in enumerate(self.tree.get_children(), 2):
                    valores = self.tree.item(item)['values']
                    for col, valor in enumerate(valores, 1):
                        ws.cell(row=row, column=col, value=valor)
                
                # Ajustar el ancho de las columnas
                for column in ws.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = min(max_length + 2, 50)
                    ws.column_dimensions[column_letter].width = adjusted_width
                
                # Guardar el archivo
                wb.save(filename)
                
                self.show_success_notification(f"Datos exportados exitosamente a {filename}")
                
            except Exception as e:
                messagebox.showerror("Error", f"Error al exportar: {e}")
    
    # -------------------------
    # Funciones de lector barcode
    # -------------------------
    def on_key_press(self, event):
        """Captura caracteres del lector tipo pistola. Cuando detecta Return procesa el barcode."""
        now = time.time()
        # Resetear buffer si hay un gap grande entre teclas
        if now - self.last_key_time > self.scanner_inter_char_delay:
            self.barcode_buffer = ""
        self.last_key_time = now
        
        # Ignorar teclas modificadoras
        if len(event.keysym) > 1 and event.keysym not in ("Return", "BackSpace"):
            return
        
        if event.keysym == "Return":
            code = self.barcode_buffer.strip()
            if code:
                self.process_barcode(code)
            self.barcode_buffer = ""
        elif event.keysym == "BackSpace":
            self.barcode_buffer = self.barcode_buffer[:-1]
        else:
            # event.char puede ser vacío para algunas teclas, proteger
            ch = event.char
            if ch:
                self.barcode_buffer += ch
    
    def process_barcode(self, code):
        """Procesa el barcode: busca mapeo, asigna trabajador y notifica."""
        # Normalizar
        code = code.strip()
        worker = self.barcode_map.get(code)
        if worker:
            self.current_worker = worker
            self.worker_label.config(text=f"{worker.get('rut')} — {worker.get('name')}")
            
            # Si hay un proceso activo, procesar la bandeja
            if self.proceso_activo:
                self.procesar_bandeja(code)
            else:
                self.show_success_notification(f"Trabajador: {worker.get('rut')}")
        else:
            # Pedir asignación si no existe
            resp = messagebox.askyesno("Barcode Desconocido",
                                       f"Código '{code}' no está asignado. ¿Desea asignarlo ahora?")
            if resp:
                self.prompt_assign_barcode(code)
    
    def prompt_assign_barcode(self, code):
        """Pide al usuario RUT y nombre para asociar con el barcode y guarda el mapeo."""
        rut = simpledialog.askstring("Asignar RUT", f"Ingrese RUT para el código {code}:")
        if not rut:
            return
        name = simpledialog.askstring("Asignar Nombre", f"Ingrese Nombre para el RUT {rut}:")
        if name is None:
            name = ""
        worker = {"rut": rut.strip(), "name": name.strip()}
        self.barcode_map[code] = worker
        self.current_worker = worker
        self.worker_label.config(text=f"{worker.get('rut')} — {worker.get('name')}")
        # Preguntar si guardar en archivo
        save = messagebox.askyesno("Guardar Mapeo", "¿Desea guardar el mapeo en archivo para uso futuro?")
        if save:
            self.save_barcode_mapping()
            self.show_success_notification("Mapeo guardado")
    
    def load_barcode_mapping(self, default_silent=False):
        """Carga mapeo barcode->rut desde archivo JSON seleccionado o por defecto."""
        if default_silent and os.path.exists(self.mapping_filename):
            try:
                with open(self.mapping_filename, 'r', encoding='utf-8') as f:
                    self.barcode_map = json.load(f)
                return
            except Exception:
                # fallo silencioso en carga por defecto
                return
        
        filename = filedialog.askopenfilename(
            title="Seleccionar archivo de mapeo (JSON)",
            filetypes=[("JSON files", "*.json"), ("All files", "*.*")]
        )
        if not filename:
            return
        try:
            with open(filename, 'r', encoding='utf-8') as f:
                self.barcode_map = json.load(f)
            self.mapping_filename = filename
            self.show_success_notification("Mapeo cargado")
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo cargar el mapeo: {e}")
    
    def save_barcode_mapping(self):
        """Guarda el mapeo actual en archivo (por defecto en home/barcode_map.json)"""
        filename = filedialog.asksaveasfilename(
            defaultextension=".json",
            initialfile=os.path.basename(self.mapping_filename) if self.mapping_filename else "barcode_map.json",
            filetypes=[("JSON files", "*.json"), ("All files", "*.*")],
            title="Guardar mapeo como..."
        )
        if not filename:
            return
        try:
            with open(filename, 'w', encoding='utf-8') as f:
                json.dump(self.barcode_map, f, indent=2, ensure_ascii=False)
            self.mapping_filename = filename
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo guardar el mapeo: {e}")
    
    # -------------------------
    # Funciones para proceso de carga
    # -------------------------
    def iniciar_proceso(self):
        """Inicia el proceso de pesaje de carga"""
        if not self.is_connected:
            messagebox.showerror("Error", "Debe estar conectado a la báscula para iniciar el proceso")
            return
        
        # Obtener peso inicial
        respuesta = messagebox.askyesno("Iniciar Proceso", 
                                       "Coloque la carga completa en la báscula y presione Sí para continuar")
        if not respuesta:
            return
        
        # Esperar un momento para estabilizar la lectura
        time.sleep(1)
        
        # Obtener el peso inicial
        self.peso_inicial = self.peso_actual
        self.peso_carga_actual = self.peso_actual
        self.peso_inicial_var.set(f"{self.peso_inicial:.2f} kg")
        self.peso_carga_actual_var.set(f"{self.peso_carga_actual:.2f} kg")
        
        # Actualizar estado
        self.proceso_activo = True
        self.estado_proceso_var.set("Activo")
        self.btn_iniciar_proceso.config(state="disabled")
        self.btn_finalizar_proceso.config(state="normal")
        
        self.show_success_notification("Proceso de carga iniciado")
    
    def finalizar_proceso(self):
        """Finaliza el proceso de pesaje de carga"""
        self.proceso_activo = False
        self.estado_proceso_var.set("Inactivo")
        self.btn_iniciar_proceso.config(state="normal")
        self.btn_finalizar_proceso.config(state="disabled")
        
        # Mostrar resumen
        self.mostrar_resumen()
        
        self.show_success_notification("Proceso de carga finalizado")
    
    def procesar_bandeja(self, codigo_barras):
        """Procesa una bandeja retirada durante el proceso de carga"""
        if not self.proceso_activo:
            messagebox.showwarning("Proceso Inactivo", "Debe iniciar un proceso de carga primero")
            return
        
        # Obtener peso antes de retirar la bandeja
        peso_antes = self.peso_carga_actual
        
        # Solicitar retirar la bandeja
        respuesta = messagebox.askyesno("Retirar Bandeja", 
                                       f"Retire la bandeja de {self.current_worker.get('name', 'el trabajador')} y presione Sí para continuar")
        if not respuesta:
            return
        
        # Esperar un momento para estabilizar la lectura
        time.sleep(1)
        
        # Obtener peso después de retirar la bandeja
        peso_despues = self.peso_actual
        self.peso_carga_actual = peso_despues
        self.peso_carga_actual_var.set(f"{peso_despues:.2f} kg")
        
        # Calcular peso de la bandeja
        peso_bandeja = peso_antes - peso_despues
        
        if peso_bandeja <= 0:
            messagebox.showerror("Error", "El peso no disminuyó después de retirar la bandeja")
            return
        
        # Registrar la bandeja retirada
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        registro = {
            'timestamp': timestamp,
            'rut': self.current_worker.get('rut', ''),
            'nombre': self.current_worker.get('name', ''),
            'peso_bandeja': peso_bandeja,
            'peso_carga_antes': peso_antes,
            'peso_carga_despues': peso_despues
        }
        
        self.bandejas_retiradas.append(registro)
        
        # Agregar al treeview
        item_count = len(self.tree.get_children())
        tag = 'evenrow' if item_count % 2 == 0 else 'oddrow'
        
        self.tree.insert("", "end", 
                        values=(timestamp, 
                                self.current_worker.get('rut', ''), 
                                self.current_worker.get('name', ''),
                                f"{peso_bandeja:.2f}",
                                f"{peso_antes:.2f}",
                                f"{peso_despues:.2f}"), 
                        tags=(tag,))
        
        # Auto-scroll al final
        children = self.tree.get_children()
        if children:
            self.tree.see(children[-1])
        
        self.show_success_notification(f"Bandeja procesada. Peso: {peso_bandeja:.2f} kg")
    
    def mostrar_resumen(self):
        """Muestra un resumen del proceso de carga"""
        if not self.bandejas_retiradas:
            messagebox.showinfo("Resumen", "No se procesaron bandejas en este ciclo")
            return
        
        # Calcular totales
        total_bandejas = len(self.bandejas_retiradas)
        total_peso = sum(b['peso_bandeja'] for b in self.bandejas_retiradas)
        
        # Agrupar por trabajador
        trabajadores = {}
        for bandeja in self.bandejas_retiradas:
            rut = bandeja['rut']
            if rut not in trabajadores:
                trabajadores[rut] = {
                    'nombre': bandeja['nombre'],
                    'cantidad': 0,
                    'peso_total': 0
                }
            trabajadores[rut]['cantidad'] += 1
            trabajadores[rut]['peso_total'] += bandeja['peso_bandeja']
        
        # Crear mensaje de resumen
        mensaje = f"RESUMEN DEL PROCESO DE CARGA\n\n"
        mensaje += f"Peso inicial: {self.peso_inicial:.2f} kg\n"
        mensaje += f"Peso final: {self.peso_carga_actual:.2f} kg\n"
        mensaje += f"Total bandejas retiradas: {total_bandejas}\n"
        mensaje += f"Peso total de bandejas: {total_peso:.2f} kg\n\n"
        
        mensaje += "DESGLOSE POR TRABAJADOR:\n"
        for rut, datos in trabajadores.items():
            mensaje += f"- {datos['nombre']} ({rut}): {datos['cantidad']} bandejas, {datos['peso_total']:.2f} kg\n"
        
        messagebox.showinfo("Resumen del Proceso", mensaje)

def main():
    root = tk.Tk()
    app = ModernPesaUSBMonitor(root)
    
    # Aplicar efectos hover
    app.add_hover_effects()
    
    def on_closing():
        app.desconectar()
        root.destroy()
    
    root.protocol("WM_DELETE_WINDOW", on_closing)
    root.mainloop()

if __name__ == "__main__":
    main()